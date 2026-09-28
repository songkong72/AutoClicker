import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_complete_clean_manual():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Minimal Modern Dark Palette
    c_bg = RGBColor(15, 23, 42)          # Deep Slate #0F172A
    c_card = RGBColor(30, 41, 59)        # Card Background #1E293B
    c_border = RGBColor(51, 65, 85)      # Border #334155
    c_primary = RGBColor(37, 99, 235)    # Brand Blue #2563EB
    c_accent = RGBColor(56, 189, 248)    # Cyan Accent #38BDF8
    c_danger = RGBColor(239, 68, 68)     # Warning Red #EF4444
    c_success = RGBColor(34, 197, 94)    # Emerald Green #22C55E
    c_warning = RGBColor(245, 158, 11)   # Amber #F59E0B
    c_white = RGBColor(255, 255, 255)
    c_muted = RGBColor(148, 163, 184)    # Muted Slate #94A3B8

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()
        return bg

    def add_header(slide, step_badge, title, subtitle):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.55), Inches(1.8), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c_primary
        badge.line.fill.background()
        tf = badge.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = step_badge
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = c_white

        tb = slide.shapes.add_textbox(Inches(2.8), Inches(0.48), Inches(9.8), Inches(0.55))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = c_white

        tb_sub = slide.shapes.add_textbox(Inches(2.8), Inches(1.05), Inches(9.8), Inches(0.35))
        tf_sub = tb_sub.text_frame
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(13.5)
        p_sub.font.color.rgb = c_muted

    # -------------------------------------------------------------
    # SLIDE 1: Cover
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_card
    card1.line.color.rgb = c_border
    card1.line.width = Pt(1.5)

    tag = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(1.6), Inches(2.3), Inches(0.45))
    tag.fill.solid()
    tag.fill.fore_color.rgb = c_primary
    tag.line.fill.background()
    p = tag.text_frame.paragraphs[0]
    p.text = "⚡ 모바일 게임 전용"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = c_white

    tb1 = s1.shapes.add_textbox(Inches(1.6), Inches(2.3), Inches(10.0), Inches(1.8))
    tf1 = tb1.text_frame
    p = tf1.paragraphs[0]
    p.text = "오토클리커 Pro (AutoClicker Pro)"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_white

    p2 = tf1.add_paragraph()
    p2.text = "초간편 APK 설치 & 완벽 정복 매뉴얼"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = c_accent
    p2.space_before = Pt(8)

    feats = [
        "🛡️ 루팅 불필요 & 원스위치 초간편 권한 (복잡한 설정 0개)",
        "⚠️ 'Play 프로텍트 차단' & '제한된 설정 허용' 완벽 해결법 수록",
        "🚀 스마트 1버튼 [띄우기/숨기기] 토글 & 100% 무한 연타 보장",
        "🛑 스마트폰 옆면 [물리 볼륨 키] 0.1초 즉시 긴급 정지"
    ]
    tb_f = s1.shapes.add_textbox(Inches(1.6), Inches(4.3), Inches(10.0), Inches(1.8))
    tf_f = tb_f.text_frame
    for i, f in enumerate(feats):
        pf = tf_f.paragraphs[0] if i == 0 else tf_f.add_paragraph()
        pf.text = f
        pf.font.size = Pt(15)
        pf.font.color.rgb = c_white
        pf.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 2: Step 1 - 파일 폰으로 옮기기
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "STEP 01", "APK 파일 스마트폰으로 옮기기", "PC에 있는 AutoClicker-Pro.apk 파일을 스마트폰으로 전송하는 3가지 방법")

    boxes_s2 = [
        ("방법 1: 카카오톡 나에게 보내기", "PC 카톡 '나와의 채팅'에 APK 파일 드래그 전송\n→ 폰 카톡에서 파일 터치하여 다운로드", c_primary),
        ("방법 2: 구글 드라이브 활용", "PC 웹 구글 드라이브에 파일 업로드\n→ 폰 구글 드라이브 앱에서 파일 옆 점3개(⋮)\n누르고 [다운로드]", c_accent),
        ("방법 3: USB 케이블 연결", "스마트폰을 PC에 USB로 연결\n→ 스마트폰 'Download' 폴더로\n파일을 직접 복사-붙여넣기", c_success)
    ]
    for i, (btitle, bdesc, bcol) in enumerate(boxes_s2):
        x = Inches(0.8 + i * 3.9)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = c_card
        c.line.color.rgb = bcol
        c.line.width = Pt(1.5)

        h = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), Inches(2.1), Inches(3.1), Inches(0.6))
        h.fill.solid()
        h.fill.fore_color.rgb = bcol
        h.line.fill.background()
        p = h.text_frame.paragraphs[0]
        p.text = btitle
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = c_white

        tb = s2.shapes.add_textbox(x + Inches(0.3), Inches(2.9), Inches(3.1), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = bdesc
        p.font.size = Pt(14)
        p.font.color.rgb = c_white
        p.line_spacing = 1.35

    # -------------------------------------------------------------
    # SLIDE 3: Play Protect 차단 완벽 해결
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "해결 01", "'Google Play 프로텍트 앱 차단됨' 30초 해결법", "스토어 외에서 설치하는 접근성 APK에 대해 구글이 띄우는 정상 보안 경고입니다.")

    c_el = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    c_el.fill.solid()
    c_el.fill.fore_color.rgb = c_card
    c_el.line.color.rgb = c_danger
    c_el.line.width = Pt(2)

    tb_el = s3.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(3.9), Inches(4.4))
    tf_el = tb_el.text_frame
    tf_el.word_wrap = True
    p = tf_el.paragraphs[0]
    p.text = "🚨 왜 차단 창이 뜰까요?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_danger

    el_descs = [
        "스마트폰의 구글 Play 프로텍트는 스토어 외에서 직접 설치하는 앱이 [접근성 권한]을 갖고 있으면 강력 차단창을 띄웁니다.",
        "",
        "💡 안심하세요!",
        "개발 중인 순수 로컬 앱이라 뜨는 것이며, 개인정보를 절대 수집하지 않는 안전한 앱입니다.",
        "오른쪽 30초 방법으로 즉시 뚫고 설치하실 수 있습니다."
    ]
    for ed in el_descs:
        p = tf_el.add_paragraph()
        p.text = ed
        p.font.size = Pt(13)
        p.font.color.rgb = c_white if not ed.startswith("💡") else c_accent
        p.space_before = Pt(4)

    c_er = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
    c_er.fill.solid()
    c_er.fill.fore_color.rgb = c_card
    c_er.line.color.rgb = c_success
    c_er.line.width = Pt(2)

    tb_er = s3.shapes.add_textbox(Inches(5.9), Inches(2.1), Inches(6.3), Inches(4.4))
    tf_er = tb_er.text_frame
    tf_er.word_wrap = True
    p = tf_er.paragraphs[0]
    p.text = "⚡ 30초 만에 뚫고 설치하는 순서 (초간단!)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_success

    sol_steps = [
        "1. 스마트폰 홈 화면에서 [Play 스토어] 앱 실행",
        "2. 우측 상단 [내 프로필 사진 (동그라미)] 터치",
        "3. 메뉴에서 [Play 프로텍트] 선택",
        "4. 화면 우측 상단 [톱니바퀴 ⚙️ (설정)] 터치",
        "5. 맨 위 'Play 프로텍트로 앱 검사' 스위치를 잠시 [꺼짐(OFF)]!",
        "6. 이제 다시 APK 파일을 누르면 차단 없이 바로 [설치] 완료!",
        "※ 설치가 끝나면 아까 끈 스위치를 다시 [켜기]로 켜두시면 됩니다."
    ]
    for ss in sol_steps:
        p = tf_er.add_paragraph()
        p.text = ss
        p.font.size = Pt(13.5)
        p.font.color.rgb = c_white if not ss.startswith("※") else c_warning
        p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 4: 제한된 설정 허용 (★ 신규 요청 핵심 슬라이드)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "해결 02", "안드로이드 '제한된 설정 허용' 푸는 법 (필수!)", "접근성 스위치가 회색으로 비활성화되어 안 켜질 때 단 10초 만에 푸는 방법")

    # Left: Problem Description
    c_s4_l = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.0))
    c_s4_l.fill.solid()
    c_s4_l.fill.fore_color.rgb = c_card
    c_s4_l.line.color.rgb = c_warning
    c_s4_l.line.width = Pt(2)

    tb_s4_l = s4.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(4.2), Inches(4.4))
    tf_s4_l = tb_s4_l.text_frame
    tf_s4_l.word_wrap = True
    p = tf_s4_l.paragraphs[0]
    p.text = "🔒 '제한된 설정' 증상"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_warning

    s4_prob = [
        "안드로이드 13 / 14 / 15 폰에서 외부 APK를 설치한 뒤 접근성을 켜려고 하면:",
        "",
        "\"보안을 위해 현재 이 설정을 사용할 수 없습니다\"라는 팝업이 뜨고 스위치가 회색으로 잠깁니다.",
        "",
        "이것은 구글의 새로운 보안 잠금이며, 오른쪽 순서대로 [제한된 설정 허용]을 1번만 눌러주시면 즉시 잠금이 풀립니다!"
    ]
    for sp in s4_prob:
        p = tf_s4_l.add_paragraph()
        p.text = sp
        p.font.size = Pt(13.5)
        p.font.color.rgb = c_white
        p.space_before = Pt(4)

    # Right: How to unlock Restricted Settings
    c_s4_r = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(1.8), Inches(6.6), Inches(5.0))
    c_s4_r.fill.solid()
    c_s4_r.fill.fore_color.rgb = c_card
    c_s4_r.line.color.rgb = c_primary
    c_s4_r.line.width = Pt(2)

    tb_s4_r = s4.shapes.add_textbox(Inches(6.2), Inches(2.1), Inches(6.0), Inches(4.4))
    tf_s4_r = tb_s4_r.text_frame
    tf_s4_r.word_wrap = True
    p = tf_s4_r.paragraphs[0]
    p.text = "🔓 제한된 설정 허용 해제 순서 (단 10초)"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent

    unlock_steps = [
        "1. 스마트폰 [설정] 앱 실행",
        "2. [애플리케이션] (또는 '앱') 메뉴로 이동",
        "3. 앱 목록에서 [오토클리커 Pro]를 찾아 터치",
        "4. 화면 우측 상단 [점 3개 (⋮)] 터치!",
        "5. [제한된 설정 허용] 터치!",
        "6. 스마트폰 화면 잠금(지문 또는 패턴/PIN) 인증",
        "🎉 이제 접근성 스위치의 잠금이 완전히 풀려서 바로 켜집니다!"
    ]
    for us in unlock_steps:
        p = tf_s4_r.add_paragraph()
        p.text = us
        p.font.size = Pt(13.5)
        p.font.color.rgb = c_white if not us.startswith("🎉") else c_success
        p.space_before = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 5: Step 2 - 원스위치 권한 설정
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "STEP 02", "평생 딱 1번! 접근성 스위치 1개만 켜기", "다른 복잡한 권한 0개! 스위치 1개만 켜면 모든 준비가 끝납니다.")

    c_s5_l = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.8), Inches(5.0))
    c_s5_l.fill.solid()
    c_s5_l.fill.fore_color.rgb = c_card
    c_s5_l.line.color.rgb = c_border

    tb_s5 = s5.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(6.2), Inches(4.4))
    tf_s5 = tb_s5.text_frame
    tf_s5.word_wrap = True
    p = tf_s5.paragraphs[0]
    p.text = "📱 설정 방법 (단 10초)"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent

    s5_steps = [
        "1. [오토클리커 Pro] 앱 실행 후 [설정하기] 터치",
        "2. 스마트폰 [접근성] 설정 화면으로 자동 이동",
        "3. [설치된 앱] → [오토클리커 Pro] 선택",
        "4. 맨 위 '오토클리커 Pro 사용' 스위치 [ON 켜기]!",
        "💡 폰을 껐다 켜도 꺼지지 않으므로 평생 딱 1번만 켜두시면 됩니다."
    ]
    for st in s5_steps:
        p = tf_s5.add_paragraph()
        p.text = st
        p.font.size = Pt(14)
        p.font.color.rgb = c_white
        p.space_before = Pt(10)

    c_s5_r = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.0))
    c_s5_r.fill.solid()
    c_s5_r.fill.fore_color.rgb = c_card
    c_s5_r.line.color.rgb = c_warning
    c_s5_r.line.width = Pt(1.5)

    tb_s5_r = s5.shapes.add_textbox(Inches(8.2), Inches(2.1), Inches(4.0), Inches(4.4))
    tf_s5_r = tb_s5_r.text_frame
    tf_s5_r.word_wrap = True
    p = tf_s5_r.paragraphs[0]
    p.text = "🎯 화면 모서리 버튼 없애기"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_warning

    tips_s5 = [
        "화면 모서리에 붙어 있는 과녁 아이콘은 안드로이드 시스템의 '바로가기' 버튼입니다.",
        "",
        "평소 화면을 가리므로 꺼두시는 게 100배 깔끔합니다!",
        "",
        "🛠️ 끄는 방법:",
        "설정에서 '오토클리커 Pro 바로가기' 스위치를 [OFF 끄기]로 두세요.",
        "오직 '사용' 스위치 1개만 켜두시면 완벽합니다."
    ]
    for tp in tips_s5:
        p = tf_s5_r.add_paragraph()
        p.text = tp
        p.font.size = Pt(13)
        p.font.color.rgb = c_white
        p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 6: Step 3 - 단 1개의 스마트 버튼으로 제어
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "STEP 03", "단 1개의 스마트 버튼으로 [띄우기 / 숨기기]", "스위치 켜도 자동 실행 X → 내가 버튼을 누를 때만 화면에 깔끔하게 나타납니다.")

    c_s6_1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c_s6_1.fill.solid()
    c_s6_1.fill.fore_color.rgb = c_card
    c_s6_1.line.color.rgb = c_primary
    c_s6_1.line.width = Pt(2)

    btn1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(4.8), Inches(0.8))
    btn1.fill.solid()
    btn1.fill.fore_color.rgb = c_primary
    btn1.line.fill.background()
    p = btn1.text_frame.paragraphs[0]
    p.text = "🚀 오토클리커 띄우기"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_white

    tb_s6_1 = s6.shapes.add_textbox(Inches(1.2), Inches(3.3), Inches(4.8), Inches(3.2))
    tf_s6_1 = tb_s6_1.text_frame
    tf_s6_1.word_wrap = True
    d1 = [
        "• 화면에 오토클리커가 없을 때 파란색 버튼 표시",
        "• 터치하는 즉시 화면에 🎯 과녁과 컨트롤 바가 나타남",
        "• 터치와 동시에 게임 화면으로 자동 즉시 전환되어 편리함"
    ]
    for d in d1:
        p = tf_s6_1.add_paragraph() if tf_s6_1.paragraphs[0].text else tf_s6_1.paragraphs[0]
        p.text = d
        p.font.size = Pt(14)
        p.font.color.rgb = c_white
        p.space_before = Pt(10)

    c_s6_2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
    c_s6_2.fill.solid()
    c_s6_2.fill.fore_color.rgb = c_card
    c_s6_2.line.color.rgb = c_danger
    c_s6_2.line.width = Pt(2)

    btn2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(2.2), Inches(4.8), Inches(0.8))
    btn2.fill.solid()
    btn2.fill.fore_color.rgb = c_danger
    btn2.line.fill.background()
    p = btn2.text_frame.paragraphs[0]
    p.text = "✕ 오토클리커 숨기기"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = c_white

    tb_s6_2 = s6.shapes.add_textbox(Inches(7.3), Inches(3.3), Inches(4.8), Inches(3.2))
    tf_s6_2 = tb_s6_2.text_frame
    tf_s6_2.word_wrap = True
    d2 = [
        "• 화면에 이미 과녁/컨트롤러가 떠 있을 때 빨간색으로 자동 전환",
        "• 터치 시 화면의 과녁과 시작 바가 즉시 사라지고 숨겨짐",
        "• 게임 화면 플로팅 바에서 [✕]를 눌러 닫아도 자동으로 [띄우기]로 복구!"
    ]
    for d in d2:
        p = tf_s6_2.add_paragraph() if tf_s6_2.paragraphs[0].text else tf_s6_2.paragraphs[0]
        p.text = d
        p.font.size = Pt(14)
        p.font.color.rgb = c_white
        p.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 7: Step 4 - 사이드 세로 플로팅 바 & 미니 타겟 사용법
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "STEP 04", "사이드 세로 캡슐 바 & 38dp 미니 과녁 실전 가이드", "화면 가장자리 슬림 바에서 시작/접기/타겟숨김 원터치 제어 & 볼륨 키 긴급 정지")

    c_s7_l = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0))
    c_s7_l.fill.solid()
    c_s7_l.fill.fore_color.rgb = c_card
    c_s7_l.line.color.rgb = c_primary
    c_s7_l.line.width = Pt(1.5)

    tb_s7_l = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.4), Inches(4.4))
    tf_s7_l = tb_s7_l.text_frame
    tf_s7_l.word_wrap = True
    p = tf_s7_l.paragraphs[0]
    p.text = "🕹️ 슬림 세로 캡슐 바 완벽 기능"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent

    s7_descs = [
        "• [ ▶ / ⏸ ]: 연타 시작(초록) 및 멈춤(빨강)",
        "• [ 👁 ]: 타겟 조준점만 쏙 숨기기/보이기 (연타는 유지!)",
        "• [ ⚙️ ]: 인게임 설정 팝업! (게임 중 속도 & 반복조건 변경)",
        "   - ♾️ 무한 반복: 수동 정지할 때까지 계속 연타",
        "   - 🔢 횟수 지정: N회(100회 등) 클릭 후 자동 정지",
        "   - ⏱️ 타이머: N분(1분 등) 경과 후 자동 정지",
        "• [ ✕ ]: 오토클리커 플로팅 바 & 타겟 완전 종료",
        "• [ ^ / v ]: 접기/펼치기 (^ 누르면 시작/정지만 앙증맞게 남음)"
    ]
    for sd in s7_descs:
        p = tf_s7_l.add_paragraph()
        p.text = sd
        p.font.size = Pt(13)
        p.font.color.rgb = c_white
        p.space_before = Pt(5)

    c_s7_r = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.0))
    c_s7_r.fill.solid()
    c_s7_r.fill.fore_color.rgb = c_card
    c_s7_r.line.color.rgb = c_danger
    c_s7_r.line.width = Pt(2)

    tb_s7_r = s7.shapes.add_textbox(Inches(7.4), Inches(2.1), Inches(4.8), Inches(4.4))
    tf_s7_r = tb_s7_r.text_frame
    tf_s7_r.word_wrap = True
    p = tf_s7_r.paragraphs[0]
    p.text = "🎯 38dp 미니 과녁 & 볼륨 키 정지"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_danger

    stop_tips = [
        "• 🎯 38dp 정밀 과녁:",
        "   화면 중앙에 앙증맞게 뜨며, 게임의 [치료] 버튼 위로 드래그 배치! 버튼 글씨를 가리지 않는 최적 크기.",
        "",
        "• 📱 스마트폰 옆면 [물리 볼륨 키] 0.1초 긴급 정지:",
        "   게임 중 화면을 터치할 필요 없이, 옆면 볼륨 UP 또는 DOWN 키를 딸깍 1번 누르면 연타가 즉시 멈춥니다!",
        "",
        "",
        "🛡️ 4중 안전 보호 시스템 완비"
    ]
    for stp in stop_tips:
        p = tf_s7_r.add_paragraph()
        p.text = stp
        p.font.size = Pt(13)
        p.font.color.rgb = c_white if not stp.startswith("• 📱") else c_accent
        p.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 8: Summary & FAQ
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "FAQ", "자주 묻는 질문 & 핵심 요약", "사용 중 궁금하신 점들을 아주 명쾌하게 정리했습니다.")

    faqs = [
        ("Q1. 연타 속도는 몇 초가 가장 좋은가요?",
         "A. 기본값인 500ms (0.5초)가 가장 좋습니다! 너무 빠르면 게임 서버가 터치를 씹을 수 있어 0.5초가 가장 안전하고 확실합니다."),
        ("Q2. 구글 플레이스토어에 정식 출시된 후에는 어떻게 설치하나요?",
         "A. 스토어 정식 출시 후에는 'Play 프로텍트 차단'이나 '제한된 설정' 해제 과정 전혀 없이, 스토어에서 검색하여 버튼 1번으로 누구나 바로 다운로드됩니다."),
        ("Q3. 플로팅 과녁과 시작 바는 어떻게 완전히 치우나요?",
         "A. 플로팅 컨트롤 바의 빨간색 [✕] 버튼을 누르거나, 앱 메인 화면의 [✕ 오토클리커 숨기기]를 누르시면 화면에서 흔적 없이 완전히 사라집니다.")
    ]
    for i, (q, a) in enumerate(faqs):
        y = Inches(1.8 + i * 1.6)
        fc = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.4))
        fc.fill.solid()
        fc.fill.fore_color.rgb = c_card
        fc.line.color.rgb = c_border

        tb = s8.shapes.add_textbox(Inches(1.1), y + Inches(0.1), Inches(11.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = q
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = c_accent

        p2 = tf.add_paragraph()
        p2.text = a
        p2.font.size = Pt(13.5)
        p2.font.color.rgb = c_white
        p2.space_before = Pt(4)

    out_path = r"E:\project\workspace\AutoClicker\AutoClicker_Pro_Manual.pptx"
    prs.save(out_path)
    print(f"Updated complete manual saved to: {out_path}")

if __name__ == "__main__":
    create_complete_clean_manual()
